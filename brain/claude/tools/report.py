#!/usr/bin/env python3
"""Charts and Excel from tables: one run turns CSV / JSON files into
  - a sparkline per series on stdout (paste it in the answer) + min / avg / max / last,
  - report.html: every chart (hover for values, light / dark) with its data table under it, one self-contained file,
  - report.xlsx: one sheet per file (header, filters, frozen top row) with a native Excel chart, plus a Summary sheet.

Input: wide tables, first column = x (time, number or name), every other column = one series of numbers.
  CSV with a header row, or JSON (a list of objects, or {"columns": [...], "rows": [[...]]}).
  Time: ISO 8601 / "YYYY-MM-DD HH:MM[:SS]" / epoch seconds or ms. The file name is the chart title
  ("aggr1 throughput MBps.csv" -> "aggr1 throughput MBps").

    python3 report.py aggr1_throughput.csv node_latency.csv --title "ONTAP last 10 days" --unit "MB/s"
    python3 report.py unused_luns.csv --kind table --title "Unused LUNs"        (Excel + HTML table, no chart)
    python3 report.py --send ~/storage-reports/x/custom.xlsx    (a file made another way: deliver it to the PC too)
Every run queues its folder for get-reports.ps1 (the PC downloads and opens it by itself); in an SSH session
it also prints the PowerShell line that copies and opens it by hand.
Options: --kind auto|line|area|bar|table (auto: line for time / numbers, bar for names), --unit U (one for all,
or one per file: --unit "MB/s,ms,%"), --out DIR
(default ~/storage-reports/<date>-<title>), --text spark|full (stdout: one sparkline per series, or the
full plotext chart; --width N columns, default 48 / 72).
Standard library + the pure-Python wheels in vendor/ (openpyxl, plotext): nothing to pip install.
"""
import argparse, csv, datetime as dt, glob, html, json, math, os, re, sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path[:0] = sorted(glob.glob(os.path.join(HERE, "vendor", "*.whl")))   # wheels import straight from the zip

# categorical slots, fixed order (validated palette: dataviz reference instance), light / dark
LIGHT = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4", "#008300", "#6250d6", "#e34948"]
DARK = ["#3987e5", "#d95926", "#199e70", "#c98500", "#d55181", "#008300", "#9085e9", "#e66767"]
MAX_SERIES = len(LIGHT)


# ---------------------------------------------------------------- reading
def read_table(path):
    if path.lower().endswith(".json"):
        d = json.load(open(path, encoding="utf-8"))
        if isinstance(d, dict) and "rows" in d:
            return [str(c) for c in d["columns"]], [list(r) for r in d["rows"]]
        cols = list(dict.fromkeys(k for row in d for k in row))
        return cols, [[row.get(c) for c in cols] for row in d]
    with open(path, newline="", encoding="utf-8-sig") as f:
        rows = list(csv.reader(f))
    return rows[0], [r for r in rows[1:] if any(c.strip() for c in r)]


def to_time(v):
    if isinstance(v, (int, float)) or re.fullmatch(r"\d{9,13}(\.\d+)?", str(v).strip()):
        x = float(v)
        return dt.datetime.fromtimestamp(x / 1000 if x > 1e11 else x)
    s = str(v).strip().replace("Z", "+00:00")
    s = re.sub(r"\.(\d+)", lambda m: "." + (m.group(1) + "000000")[:6], s, count=1)   # Python 3.9: exactly 6 digits
    t = dt.datetime.fromisoformat(s if "T" in s or len(s) <= 10 else s.replace(" ", "T", 1))
    return t.astimezone().replace(tzinfo=None) if t.tzinfo else t


def to_num(v):
    if v is None or str(v).strip() in ("", "-", "null", "None", "NaN", "nan"):
        return None
    return float(str(v).replace(",", "").strip())


def load(path):
    cols, rows = read_table(path)
    xs = [r[0] for r in rows]
    try:
        xs, xtype = [to_time(x) for x in xs], "time"
    except (ValueError, TypeError, OverflowError, OSError):
        try:
            xs, xtype = [float(x) for x in xs], "number"
        except (ValueError, TypeError):
            xs, xtype = [str(x) for x in xs], "name"
    series = {}
    for i, name in enumerate(cols[1:], 1):
        try:
            series[str(name)] = [to_num(r[i] if i < len(r) else None) for r in rows]
        except ValueError:
            series[str(name)] = None                       # a text column: kept in the table, not charted
    if xtype == "time":                                    # sorted by time
        order = sorted(range(len(xs)), key=lambda k: xs[k])
        xs = [xs[k] for k in order]
        rows = [rows[k] for k in order]
        series = {n: (None if v is None else [v[k] for k in order]) for n, v in series.items()}
    title = re.sub(r"[_]+", " ", os.path.splitext(os.path.basename(path))[0]).strip()
    return {"title": title, "cols": cols, "rows": rows, "x": xs, "xtype": xtype,
            "series": {n: v for n, v in series.items() if v is not None}}


def limit_series(t):
    """At most 8 series (one per palette slot, never a generated hue): the 8 with the highest peak."""
    s = t["series"]
    if len(s) <= MAX_SERIES:
        return []
    keep = sorted(s, key=lambda n: -max((v for v in s[n] if v is not None), default=0))[:MAX_SERIES]
    dropped = [n for n in s if n not in keep]
    t["series"] = {n: s[n] for n in s if n in keep}
    return dropped


# ---------------------------------------------------------------- numbers
def fmt(v, unit=""):
    if v is None:
        return "-"
    a = abs(v)
    s = f"{v:,.0f}" if a >= 100 else f"{v:,.1f}" if a >= 10 else f"{v:,.2f}" if a >= 0.01 or a == 0 else f"{v:.2e}"
    return f"{s} {unit}".strip()


def nice_ticks(lo, hi, n=5):
    if hi <= lo:
        hi = lo + 1
    step = 10 ** math.floor(math.log10((hi - lo) / n))
    for m in (1, 2, 2.5, 5, 10):
        if (hi - lo) / (step * m) <= n:
            step *= m
            break
    start = math.floor(lo / step) * step
    return [start + i * step for i in range(int(math.ceil((hi - start) / step)) + 1)]


def tick_fmt(v, ticks):
    """Axis labels share one number of decimals, from the tick step."""
    step = abs(ticks[1] - ticks[0]) if len(ticks) > 1 else 1
    d = 0 if step >= 1 else min(int(-math.floor(math.log10(step))), 6)
    return f"{v:,.{d}f}"


def xlabel(x, xtype, span_days):
    if xtype != "time":
        return fmt(x) if xtype == "number" else str(x)
    return x.strftime("%b %d") if span_days > 2 else x.strftime("%b %d %H:%M") if span_days > 0.5 else x.strftime("%H:%M")


def stats(t, unit):
    out = []
    for name, v in t["series"].items():
        vals = [(x, y) for x, y in zip(t["x"], v) if y is not None]
        if not vals:
            continue
        peak = max(vals, key=lambda p: p[1])
        out.append({"series": name, "min": min(y for _, y in vals), "avg": sum(y for _, y in vals) / len(vals),
                    "max": peak[1], "max_at": peak[0], "last": vals[-1][1], "points": len(vals)})
    return out


# ---------------------------------------------------------------- text chart (stdout)
def spark(v, n=48):
    """One line of blocks (bucket max, n wide); gaps are blanks."""
    b = [[y for y in v[i * len(v) // n:(i + 1) * len(v) // n] if y is not None] for i in range(n)] if len(v) > n \
        else [[y] if y is not None else [] for y in v]
    b = [max(x) if x else None for x in b]
    vals = [x for x in b if x is not None]
    lo, hi = min(vals, default=0), max(vals, default=0)
    return "".join(" " if x is None else "▄" if hi == lo else "▁▂▃▄▅▆▇█"[int((x - lo) / (hi - lo) * 7 + .5)] for x in b)


def text_chart(t, kind, unit, width):
    """--text full: one plotext chart per series; the answer shows no colour, so series are told apart by titles."""
    import plotext as plt
    out, names = [], list(t["series"])
    span = (t["x"][-1] - t["x"][0]).total_seconds() / 86400 if t["xtype"] == "time" and len(t["x"]) > 1 else 0
    for name in names:
        v = t["series"][name]
        plt.clf()
        plt.plotsize(width, min(len(v), 20) + 4 if kind == "bar" else 10)
        plt.theme("clear")
        label = (name if len(names) > 1 else t["title"]) + (f" ({unit})" if unit else "")
        if kind == "bar":
            plt.bar([str(x) for x in t["x"]][:20], [y or 0 for y in v][:20], orientation="horizontal", width=0.4)
        else:
            pts = [(i, y) for i, y in enumerate(v) if y is not None]
            plt.plot([p[0] for p in pts], [p[1] for p in pts], marker="braille")
            ticks = sorted(set(round(i * (len(v) - 1) / 5) for i in range(6))) if len(v) > 1 else [0]
            plt.xticks(ticks, [xlabel(t["x"][i], t["xtype"], span) for i in ticks])
        plt.title(label)
        out.append(re.sub(r"\x1b\[[0-9;]*m", "", plt.build()))
    return "\n".join(out)


# ---------------------------------------------------------------- HTML (SVG charts, hover, table)
W, H, ML, MR, MT, MB = 960, 340, 64, 16, 16, 40


def svg_chart(t, kind, unit, cid):
    names = list(t["series"])
    allv = [y for v in t["series"].values() for y in v if y is not None] or [0]
    lo, hi = min(0, min(allv)), max(allv)
    ticks = nice_ticks(lo, hi)
    lo, hi = ticks[0], ticks[-1]
    pw, ph = W - ML - MR, H - MT - MB
    ys = lambda v: MT + ph - (v - lo) / (hi - lo) * ph
    n = len(t["x"])
    if t["xtype"] == "time" and n > 1:
        t0, t1 = t["x"][0].timestamp(), t["x"][-1].timestamp()
        xs = [ML + (x.timestamp() - t0) / max(t1 - t0, 1) * pw for x in t["x"]]
        span = (t1 - t0) / 86400
    elif t["xtype"] == "number" and n > 1:
        xs = [ML + (x - t["x"][0]) / max(t["x"][-1] - t["x"][0], 1e-9) * pw for x in t["x"]]
        span = 0
    else:
        xs = [ML + (i + 0.5) * pw / max(n, 1) for i in range(n)]
        span = 0
    p = [f'<svg viewBox="0 0 {W} {H}" role="img" aria-label="{html.escape(t["title"])}" data-chart="{cid}">']
    for tk in ticks:
        y = ys(tk)
        p.append(f'<line class="grid{" base" if tk == 0 else ""}" x1="{ML}" x2="{W - MR}" y1="{y:.1f}" y2="{y:.1f}"/>'
                 f'<text class="tick" x="{ML - 8}" y="{y + 4:.1f}" text-anchor="end">{tick_fmt(tk, ticks)}</text>')
    step = max(1, n // 7) if kind == "bar" else max(1, (n - 1) // 6)
    for i in range(0, n, step):
        p.append(f'<text class="tick" x="{xs[i]:.1f}" y="{H - MB + 18}" text-anchor="middle">'
                 f'{html.escape(xlabel(t["x"][i], t["xtype"], span))}</text>')
    if kind == "bar":
        gw = pw / max(n, 1) * 0.8
        bw = max(gw / len(names) - 2, 1)                   # 2px surface gap between bars
        for s, name in enumerate(names):
            for i, v in enumerate(t["series"][name]):
                if v is None:
                    continue
                x = xs[i] - gw / 2 + s * (bw + 2)
                y0, y1 = sorted((ys(0), ys(v)))
                p.append(f'<rect class="s{s}" x="{x:.1f}" y="{y0:.1f}" width="{bw:.1f}" height="{max(y1 - y0, 1):.1f}" rx="2"/>')
    else:
        for s, name in enumerate(names):
            segs, cur = [], []
            for x, v in zip(xs, t["series"][name]):
                if v is None:
                    if cur:
                        segs.append(cur)
                    cur = []
                else:
                    cur.append(f"{x:.1f},{ys(v):.1f}")
            if cur:
                segs.append(cur)
            for seg in segs:
                if kind == "area":
                    base = f"{seg[-1].split(',')[0]},{ys(max(lo, 0)):.1f} {seg[0].split(',')[0]},{ys(max(lo, 0)):.1f}"
                    p.append(f'<polygon class="a{s}" points="{" ".join(seg)} {base}"/>')
                p.append(f'<polyline class="s{s}" points="{" ".join(seg)}"/>')
    p.append(f'<line class="cross" x1="0" x2="0" y1="{MT}" y2="{MT + ph}" visibility="hidden"/>')
    p.append(f'<rect class="hit" x="{ML}" y="{MT}" width="{pw}" height="{ph}"/></svg>')
    data = {"x": [x.strftime("%a %b %d %H:%M") if t["xtype"] == "time" else (fmt(x) if t["xtype"] == "number" else x)
                  for x in t["x"]],
            "px": [round(x, 1) for x in xs], "s": names, "v": [t["series"][nm] for nm in names], "unit": unit}
    return "".join(p), data


def table_html(t, limit=2000):
    head = "".join(f"<th>{html.escape(str(c))}</th>" for c in t["cols"])
    body = "".join("<tr>" + "".join(f"<td>{html.escape(str(c))}</td>" for c in r) + "</tr>" for r in t["rows"][:limit])
    more = f"<p class='muted'>first {limit} of {len(t['rows'])} rows (all rows: report.xlsx)</p>" if len(t["rows"]) > limit else ""
    return f"<div class='tbl'><table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table></div>{more}"


CSS = """
.viz-root{color-scheme:light;--surface-1:#fcfcfb;--surface-2:#f3f2ef;--text-primary:#0b0b0b;--text-secondary:#52514e;
 --grid:#e4e3df;--axis:#a3a29c;%s}
@media (prefers-color-scheme:dark){:root:where(:not([data-theme="light"])) .viz-root{color-scheme:dark;--surface-1:#1a1a19;
 --surface-2:#242423;--text-primary:#f0efec;--text-secondary:#c3c2b7;--grid:#333331;--axis:#6b6a64;%s}}
:root[data-theme="dark"] .viz-root{color-scheme:dark;--surface-1:#1a1a19;--surface-2:#242423;--text-primary:#f0efec;
 --text-secondary:#c3c2b7;--grid:#333331;--axis:#6b6a64;%s}
body{margin:0;background:var(--surface-1)}
.viz-root{background:var(--surface-1);color:var(--text-primary);font:14px/1.45 system-ui,-apple-system,"Segoe UI",sans-serif;
 max-width:1000px;margin:0 auto;padding:24px 16px}
h1{font-size:22px;margin:0 0 4px}h2{font-size:16px;margin:32px 0 6px}.muted{color:var(--text-secondary);font-size:12px}
.legend{display:flex;flex-wrap:wrap;gap:14px;margin:6px 0;color:var(--text-secondary);font-size:13px}
.legend i{display:inline-block;width:14px;height:3px;border-radius:2px;margin-right:6px;vertical-align:middle}
.stats{display:flex;flex-wrap:wrap;gap:8px 20px;font-size:13px;color:var(--text-secondary)}.stats b{color:var(--text-primary)}
.chart{position:relative}svg{width:100%%;height:auto;display:block}
.grid{stroke:var(--grid);stroke-width:1}.grid.base{stroke:var(--axis)}.tick{fill:var(--text-secondary);font-size:11px}
polyline{fill:none;stroke-width:2;stroke-linejoin:round;stroke-linecap:round}.cross{stroke:var(--axis);stroke-width:1}
.hit{fill:transparent}.tip{position:absolute;pointer-events:none;background:var(--surface-2);border:1px solid var(--grid);
 border-radius:6px;padding:6px 8px;font-size:12px;white-space:nowrap;display:none;box-shadow:0 2px 8px rgba(0,0,0,.15)}
.tip i{display:inline-block;width:10px;height:3px;border-radius:2px;margin-right:6px;vertical-align:middle}
details{margin-top:6px}summary{cursor:pointer;color:var(--text-secondary);font-size:13px}
.tbl{max-height:360px;overflow:auto;border:1px solid var(--grid);border-radius:6px;margin-top:6px}
table{border-collapse:collapse;font-size:12px;width:100%%}th,td{padding:4px 8px;text-align:left;border-bottom:1px solid var(--grid);
 white-space:nowrap}th{position:sticky;top:0;background:var(--surface-2)}
"""
JS = """
document.querySelectorAll('svg[data-chart]').forEach(svg=>{
 const d=DATA[svg.dataset.chart],box=svg.parentNode,tip=box.querySelector('.tip'),cr=svg.querySelector('.cross');
 const hit=svg.querySelector('.hit');if(!d.x.length)return;
 hit.addEventListener('mousemove',e=>{const r=svg.getBoundingClientRect(),k=960/r.width,mx=(e.clientX-r.left)*k;
  let i=0,b=1e9;d.px.forEach((p,j)=>{const q=Math.abs(p-mx);if(q<b){b=q;i=j}});
  cr.setAttribute('x1',d.px[i]);cr.setAttribute('x2',d.px[i]);cr.setAttribute('visibility','visible');
  tip.innerHTML='<b>'+d.x[i]+'</b>'+d.s.map((n,s)=>'<br><i style="background:var(--series-'+(s+1)+')"></i>'+n+': '+
   (d.v[s][i]==null?'-':d.v[s][i].toLocaleString(undefined,{maximumFractionDigits:2})+(d.unit?' '+d.unit:''))).join('');
  tip.style.display='block';const left=d.px[i]/k;tip.style.left=(left>r.width*0.65?left-tip.offsetWidth-12:left+12)+'px';
  tip.style.top=(e.clientY-r.top)+'px'});
 hit.addEventListener('mouseleave',()=>{tip.style.display='none';cr.setAttribute('visibility','hidden')});
});
"""


def write_html(path, title, tables, kinds, stats_by):
    var = lambda pal: "".join(f"--series-{i + 1}:{c};" for i, c in enumerate(pal))
    cls = "".join(f".s{i}{{stroke:var(--series-{i + 1});fill:var(--series-{i + 1})}}polyline.s{i}{{fill:none}}"
                  f".a{i}{{fill:var(--series-{i + 1});opacity:.15}}" for i in range(MAX_SERIES))
    parts, data = [], {}
    for n, (t, kind) in enumerate(zip(tables, kinds)):
        unit = t["unit"]
        parts.append(f"<h2>{html.escape(t['title'])}</h2>")
        if kind != "table" and t["series"]:
            svg, data[f"c{n}"] = svg_chart(t, kind, unit, f"c{n}")
            st = stats_by[n]
            parts.append("<div class='stats'>" + "".join(
                f"<span>{html.escape(s['series'])}: avg <b>{fmt(s['avg'], unit)}</b> · max <b>{fmt(s['max'], unit)}</b>"
                f" ({xlabel(s['max_at'], t['xtype'], 1)}) · last <b>{fmt(s['last'], unit)}</b></span>" for s in st) + "</div>")
            if len(t["series"]) > 1:
                parts.append("<div class='legend'>" + "".join(
                    f"<span><i style='background:var(--series-{i + 1})'></i>{html.escape(nm)}</span>"
                    for i, nm in enumerate(t["series"])) + "</div>")
            parts.append(f"<div class='chart'>{svg}<div class='tip'></div></div>")
            parts.append(f"<details><summary>Data table ({len(t['rows'])} rows)</summary>{table_html(t)}</details>")
        else:
            parts.append(table_html(t))
    page = (f"<!doctype html><html><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'>"
            f"<title>{html.escape(title)}</title><style>{CSS % (var(LIGHT), var(DARK), var(DARK))}{cls}</style></head>"
            f"<body><div class='viz-root'><h1>{html.escape(title)}</h1><p class='muted'>Generated "
            f"{dt.datetime.now():%Y-%m-%d %H:%M} · report.xlsx has the same data with Excel charts</p>{''.join(parts)}</div>"
            f"<script>const DATA={json.dumps(data, default=str)};{JS}</script></body></html>")
    open(path, "w", encoding="utf-8").write(page)


# ---------------------------------------------------------------- Excel
def write_xlsx(path, tables, kinds, stats_by):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, AreaChart, Reference
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    summ = wb.active
    summ.title = "Summary"
    summ.append(["Chart", "Series", "Min", "Avg", "Max", "Max at", "Last", "Points", "Unit"])
    used = {"Summary"}
    for t, kind, st in zip(tables, kinds, stats_by):
        unit = t["unit"]
        name = re.sub(r"[\[\]:*?/\\]", " ", t["title"])[:31] or "data"
        base, k = name, 2
        while name in used:
            name = f"{base[:28]} {k}"
            k += 1
        used.add(name)
        ws = wb.create_sheet(name)
        ws.append(t["cols"])
        numcols = {i for i, c in enumerate(t["cols"]) if str(c) in t["series"]}
        for x, r in zip(t["x"], t["rows"]):
            row = [x if t["xtype"] != "name" else r[0]]
            for i, c in enumerate(r[1:], 1):
                row.append(to_num(c) if i in numcols else c)
            ws.append(row)
        for c in ws[1]:
            c.font, c.fill = Font(bold=True), PatternFill("solid", fgColor="E4E3DF")
        ws.freeze_panes, ws.auto_filter.ref = "A2", ws.dimensions
        for i, c in enumerate(t["cols"], 1):
            width = max([len(str(c))] + [len(str(r[i - 1])) for r in t["rows"][:200] if i - 1 < len(r)])
            ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 48)
        if t["xtype"] == "time":
            for (cell,) in ws.iter_rows(min_row=2, max_col=1):
                cell.number_format = "yyyy-mm-dd hh:mm"
        if kind != "table" and t["series"]:
            ch = {"bar": BarChart, "area": AreaChart}.get(kind, LineChart)()
            ch.title, ch.height, ch.width = t["title"], 9, 22
            ch.y_axis.title = unit or None
            n = len(t["rows"]) + 1
            for i, c in enumerate(t["cols"]):
                if str(c) in t["series"]:
                    ch.add_data(Reference(ws, min_col=i + 1, min_row=1, max_row=n), titles_from_data=True)
            ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=n))
            for s, ser in enumerate(ch.series):
                color = LIGHT[s % MAX_SERIES][1:]
                if kind == "bar":
                    ser.graphicalProperties.solidFill = color
                else:
                    ser.graphicalProperties.line.solidFill = color
                    ser.graphicalProperties.line.width = 19050      # 1.5 pt
                    ser.smooth = False
            if t["xtype"] == "time":
                ch.x_axis.number_format = "mmm dd" if len(t["rows"]) > 48 else "hh:mm"
            ws.add_chart(ch, f"{get_column_letter(len(t['cols']) + 2)}2")
        for s in st:
            summ.append([t["title"], s["series"], s["min"], s["avg"], s["max"], s["max_at"], s["last"], s["points"], unit])
    for c in summ[1]:
        c.font, c.fill = Font(bold=True), PatternFill("solid", fgColor="E4E3DF")
    for col, w in zip("ABCDEFGHI", (32, 28, 12, 12, 12, 18, 12, 8, 8)):
        summ.column_dimensions[col].width = w
    for row in summ.iter_rows(min_row=2):
        for c in row[2:5] + (row[6],):
            c.number_format = "#,##0.00"
        if isinstance(row[5].value, dt.datetime):
            row[5].number_format = "yyyy-mm-dd hh:mm"
    if summ.max_row == 1:
        wb.remove(summ)
    wb.save(path)


# ---------------------------------------------------------------- to the PC
OUTBOX = os.path.expanduser("~/storage-reports/.outbox")     # read and emptied by get-reports.ps1 on the PC


def send(folder, open_name):
    """Queue the folder for get-reports.ps1 (it downloads it into Downloads\\storage-reports and opens open_name)
    and return the PowerShell line that does the same by hand. Claude runs inside his SSH session, so
    SSH_CONNECTION ("client port server port") holds the address and port he connected to."""
    import getpass, socket
    folder = os.path.abspath(os.path.expanduser(folder))
    root = os.path.dirname(OUTBOX)
    if folder == root:
        sys.exit(f"save the file in its own subfolder of {root}, not in {root} itself (the PC copies the whole folder)")
    if not re.fullmatch(r"[\w.-]+\.(html|xlsx|csv|pdf|png|svg|md|txt)", open_name, re.I):
        open_name = ""                                       # the PC only auto-opens these types
    os.makedirs(root, exist_ok=True)
    line = f"{folder}\t{open_name}"
    old = open(OUTBOX, encoding="utf-8").read().splitlines() if os.path.exists(OUTBOX) else []
    if line not in old:
        with open(OUTBOX, "a", encoding="utf-8") as f:
            f.write(line + "\n")
    parts = os.environ.get("SSH_CONNECTION", "").split()
    host = parts[2] if len(parts) == 4 else socket.getfqdn()
    port = f"-P {parts[3]} " if len(parts) == 4 and parts[3] != "22" else ""
    host = f"[{host}]" if ":" in host else host
    local = f"$env:USERPROFILE\\Downloads\\{os.path.basename(folder)}"
    return (f'scp {port}-r {getpass.getuser()}@{host}:"{folder}" "$env:USERPROFILE\\Downloads\\"' +
            (f'; ii "{local}\\{open_name}"' if open_name else ""))


# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    if sys.argv[1:2] == ["--send"] and len(sys.argv) == 3:
        f = os.path.abspath(os.path.expanduser(sys.argv[2]))
        line = send(os.path.dirname(f), os.path.basename(f))
        print(f"to the PC (PowerShell: copies into Downloads and opens it): {line}" if os.environ.get("SSH_CONNECTION")
              else f"queued {f} for the PC")
        return
    ap.add_argument("files", nargs="+")
    ap.add_argument("--title", default="")
    ap.add_argument("--unit", default="")
    ap.add_argument("--kind", default="auto", choices=["auto", "line", "area", "bar", "table"])
    ap.add_argument("--out", default="")
    ap.add_argument("--text", default="spark", choices=["spark", "full"])
    ap.add_argument("--width", type=int, default=0)
    a = ap.parse_args()
    tables = [load(f) for f in a.files]
    title = a.title or tables[0]["title"]
    out = os.path.expanduser(a.out or os.path.join("~/storage-reports", f"{dt.date.today()}-" +
                                                   re.sub(r"[^\w.-]+", "-", title.lower()).strip("-")))
    os.makedirs(out, exist_ok=True)
    kinds, stats_by = [], []
    units = [u.strip() for u in a.unit.split(",")] if "," in a.unit else [a.unit] * len(tables)
    units += [units[-1] if units else ""] * (len(tables) - len(units))
    for t, unit in zip(tables, units):
        t["unit"] = unit
        kind = a.kind if a.kind != "auto" else ("bar" if t["xtype"] == "name" else "line")
        if not t["series"]:
            kind = "table"
        kinds.append(kind)
        dropped = limit_series(t) if kind != "table" else []
        st = stats(t, unit) if kind != "table" else []
        stats_by.append(st)
        rng = f"  {xlabel(t['x'][0], t['xtype'], 0)} -> {xlabel(t['x'][-1], t['xtype'], 0)}" if kind == "line" and t["x"] else ""
        print(f"## {t['title']}  ({len(t['rows'])} rows){rng}")
        if dropped:
            print(f"note: {len(dropped)} more series not charted (8 at most; split into several files): {', '.join(dropped[:10])}")
        if kind != "table":
            if a.text == "full":
                print(text_chart(t, kind, unit, a.width or 72))
            for s in st:
                if a.text == "spark":
                    v = t["series"][s["series"]]
                    print(f"  {s['series']}: " + (" · ".join(f"{x} {fmt(y)}" for x, y in list(zip(t["x"], v))[:20] if y is not None)
                                                  if kind == "bar" else spark(v, a.width or 48)))
                at = xlabel(s["max_at"], t["xtype"], 1)
                print(f"  {s['series']}: min {fmt(s['min'], unit)} · avg {fmt(s['avg'], unit)} · "
                      f"max {fmt(s['max'], unit)} at {at} · last {fmt(s['last'], unit)}")
        print()
    write_html(os.path.join(out, "report.html"), title, tables, kinds, stats_by)
    write_xlsx(os.path.join(out, "report.xlsx"), tables, kinds, stats_by)
    first = "report.html" if any(k != "table" for k in kinds) else "report.xlsx"
    line = send(out, first)                                  # always queued for get-reports.ps1
    if os.environ.get("SSH_CONNECTION"):
        print(f"to the PC (PowerShell: copies into Downloads and opens it): {line}")
    else:
        print(f"files: {out}/report.html (charts, hover for values) · {out}/report.xlsx (data + Excel charts)")


if __name__ == "__main__":
    main()
