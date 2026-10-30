"""Check of the chart / Excel tool on tiny tables: python3 brain/test_report.py"""
import glob, json, os, subprocess, sys, tempfile

here = os.path.dirname(os.path.abspath(__file__))
tool = os.path.join(here, "claude", "tools", "report.py")
d = tempfile.mkdtemp()
os.environ["HOME"] = d                                           # the PC queue (~/storage-reports/.outbox) goes here
os.environ["SSH_CONNECTION"] = "10.0.0.9 50000 10.0.0.5 2222"
open(os.path.join(d, "vol1 throughput.csv"), "w").write(        # unsorted, a gap, two series
    "time,read,write\n2026-10-23T02:00:00,30,10\n2026-10-23T00:00:00,10,5\n2026-10-23T01:00:00,,7\n")
json.dump([{"pool": "SRP_1", "used %": 81.5}, {"pool": "SRP_2", "used %": 40}], open(os.path.join(d, "srp used.json"), "w"))
open(os.path.join(d, "unused luns.csv"), "w").write("naa,size GB\nnaa.600a0980383030,500\n")

out = subprocess.run([sys.executable, "-I", tool, os.path.join(d, "vol1 throughput.csv"), os.path.join(d, "srp used.json"),
                      "--unit", "MB/s,%", "--out", d, "--width", "60"], capture_output=True, text=True)
assert out.returncode == 0, out.stderr
o = out.stdout
assert "read: min 10.0 MB/s · avg 20.0 MB/s · max 30.0 MB/s at Oct 23 02:00 · last 30.0 MB/s" in o, o
assert "used %: min 40.0 % · avg 60.8 % · max 81.5 % at SRP_1" in o and "SRP_2" in o, o
assert '-P 2222 -r ' in o and "@10.0.0.5:" in o and 'report.html"' in o, o      # copy + open line
outbox = os.path.join(d, "storage-reports", ".outbox")
assert open(outbox).read() == d + "\treport.html\n"
assert "▁" in o or "▄" in o                                      # one sparkline per series, no plotext chart
assert "\x1b[" not in o                                          # no colour codes in the answer text
page = open(os.path.join(d, "report.html")).read()
assert page.count("<svg") == 2 and "--series-1:#2a78d6" in page and "vol1 throughput" in page

sys.path[:0] = glob.glob(os.path.join(here, "claude", "tools", "vendor", "*.whl"))
import openpyxl  # noqa: E402
wb = openpyxl.load_workbook(os.path.join(d, "report.xlsx"))
assert wb.sheetnames == ["Summary", "vol1 throughput", "srp used"], wb.sheetnames
ws = wb["vol1 throughput"]
assert ws["A2"].value.hour == 0 and ws["B3"].value is None and len(ws._charts) == 1   # sorted by time, gap kept
assert wb["Summary"].max_row == 4                                # read, write, used %

out = subprocess.run([sys.executable, "-I", tool, os.path.join(d, "unused luns.csv"), "--kind", "table", "--out", d],
                     capture_output=True, text=True)
assert out.returncode == 0 and "unused luns" in out.stdout, out.stderr
wb = openpyxl.load_workbook(os.path.join(d, "report.xlsx"))
assert wb.sheetnames == ["unused luns"] and wb["unused luns"]["B2"].value == 500

open(os.path.join(d, "frac.csv"), "w").write(                   # 2-digit and 9-digit fractions: Python 3.9 fromisoformat rejects both
    "time,v\n2026-10-23T14:05:01.123456789+00:00,2\n2026-10-23T14:05:00.12Z,1\n")
out = subprocess.run([sys.executable, "-I", tool, os.path.join(d, "frac.csv"), "--out", os.path.join(d, "f")], capture_output=True, text=True)
assert out.returncode == 0 and " -> " in out.stdout, out.stdout + out.stderr   # a time axis, not bar names
subprocess.run([sys.executable, "-I", tool, "--send", os.path.join(d, "f", "report.html")], capture_output=True, text=True)
subprocess.run([sys.executable, "-I", tool, "--send", os.path.join(d, "f", "report.html")], capture_output=True, text=True)
subprocess.run([sys.executable, "-I", tool, "--send", os.path.join(d, "f", "run.bat")], capture_output=True, text=True)
subprocess.run([sys.executable, "-I", tool, "--send", os.path.join(d, "storage-reports", "x.xlsx")], capture_output=True, text=True)
lines = open(outbox).read().splitlines()
assert lines.count(os.path.join(d, "f") + "\treport.html") == 1, lines          # queued once
assert os.path.join(d, "f") + "\t" in lines, lines                              # .bat is queued but never auto-opened
assert not any(l.startswith(os.path.join(d, "storage-reports") + "\t") for l in lines), lines   # never the whole reports folder
print("report ok")
